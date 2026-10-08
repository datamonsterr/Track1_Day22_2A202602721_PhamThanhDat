"""Use temporary LibreOffice container; transplant results, never its formula edits."""
import argparse
import json
import subprocess
import tempfile
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
CONTAINER = 'day22-office-bootstrap'


def run(*args):
    subprocess.run(args, check=True, capture_output=True, text=True)


def convert(path, extension):
    run('docker', 'exec', CONTAINER, 'mkdir', '-p', '/tmp/day22-input', '/tmp/day22-output')
    run('docker', 'cp', str(path), f'{CONTAINER}:/tmp/day22-input/{path.name}')
    output_name = path.with_suffix('.' + extension).name
    run('docker', 'exec', CONTAINER, 'rm', '-f', '/tmp/day22-output/' + output_name)
    run('docker', 'exec', CONTAINER, 'soffice', '-env:UserInstallation=file:///tmp/day22-office-profile',
        '--headless', '--convert-to', extension, '--outdir', '/tmp/day22-output',
        '/tmp/day22-input/' + path.name)
    folder = tempfile.TemporaryDirectory(prefix='day22-render-')
    output = Path(folder.name) / output_name
    run('docker', 'cp', f'{CONTAINER}:/tmp/day22-output/{output_name}', str(output))
    return folder, output


def recalculate():
    path = ROOT / 'deliverables/Day22-DevPulse-Monetization-Model.xlsx'
    folder, calculated = convert(path, 'xlsx')
    original = load_workbook(path)
    cached = load_workbook(calculated, data_only=True)
    results = {}
    count = 0
    with ZipFile(path) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    for index, sheet in enumerate(original, start=1):
        xml_name = f'xl/worksheets/sheet{index}.xml'
        root = ET.fromstring(parts[xml_name])
        for cell in root.findall('.//m:c', NS):
            if cell.find('m:f', NS) is None:
                continue
            address = cell.attrib['r']
            value = cached[sheet.title][address].value
            if value is None or cached[sheet.title][address].data_type == 'e':
                raise ValueError(f'Formula calculation failed: {sheet.title}!{address} = {value}')
            old = cell.find('m:v', NS)
            if old is not None:
                cell.remove(old)
            v = ET.SubElement(cell, '{' + NS['m'] + '}v')
            if isinstance(value, bool):
                cell.set('t', 'b')
                v.text = '1' if value else '0'
            elif isinstance(value, str):
                cell.set('t', 'str')
                v.text = value
            else:
                cell.attrib.pop('t', None)
                v.text = str(value)
            results[f'{sheet.title}!{address}'] = value
            count += 1
        parts[xml_name] = ET.tostring(root, encoding='utf-8', xml_declaration=True)
    temporary = path.with_suffix('.tmp')
    with ZipFile(temporary, 'w', ZIP_DEFLATED) as target:
        for name, data in parts.items():
            target.writestr(name, data)
    temporary.replace(path)
    (ROOT / 'data/calculated-results.json').write_text(json.dumps(results, ensure_ascii=False, indent=2) + '\n')
    folder.cleanup()
    print(f'LibreOffice evaluated {count} formulas; original formulas retained.')


def render():
    path = ROOT / 'deliverables/Day22-DevPulse-Monetization-One-Pager.docx'
    folder, pdf = convert(path, 'pdf')
    path.with_suffix('.pdf').write_bytes(pdf.read_bytes())
    folder.cleanup()
    print('Rendered DOCX to PDF with LibreOffice.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('mode', choices=['xlsx', 'pdf'])
    args = parser.parse_args()
    recalculate() if args.mode == 'xlsx' else render()
