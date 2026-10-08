"""Create compact, three-block submission from the supplied DOCX package and XLSX cache."""
import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables/Day22-DevPulse-Monetization-One-Pager.docx'
book = load_workbook(ROOT / 'deliverables/Day22-DevPulse-Monetization-Model.xlsx', data_only=True)
trace = []
keys = {book['7_Assumptions'].cell(row, 1).value: f'B{row}' for row in range(4, book['7_Assumptions'].max_row + 1)}


def value(sheet, cell, fmt=str, label=None):
    raw = book[sheet][cell].value
    if raw is None:
        raise ValueError(f'Missing calculated value: {sheet}!{cell}')
    rendered = fmt(raw)
    trace.append({'label': label or f'{sheet}!{cell}', 'sheet': sheet, 'cell': cell,
                  'value': raw, 'formatted': rendered})
    return rendered


def a(key, fmt=str):
    return value('7_Assumptions', keys[key], fmt, key)


usd = lambda x: f'${x:,.2f}'
pct = lambda x: f'{x:.2%}'
whole = lambda x: f'{x:g}'
doc = Document(ROOT / 'templates/Day22-AI-Product-GTM-One-Pager-Template.docx')
for child in list(doc.element.body):
    if child.tag != qn('w:sectPr'):
        doc.element.body.remove(child)
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(0.9)
section.left_margin = section.right_margin = Cm(1.05)
section.header_distance = section.footer_distance = Cm(0.3)
style = doc.styles['normal']
style.font.name = 'DejaVu Sans'
style.font.size = Pt(9.5)
style.paragraph_format.space_before = Pt(0)
style.paragraph_format.space_after = Pt(2)
style.paragraph_format.line_spacing = 1.0


def paragraph(text, bold=False, size=None):
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size + (1 if size < 9 else 0))
    return p


def heading(text):
    p = paragraph(text, True, 11)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.runs[0].font.color.rgb = RGBColor.from_string('17365D')


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'TableNormal'
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement('w:' + edge)
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:color'), 'AAB7C4')
        borders.append(border)
    t._tbl.tblPr.append(borders)
    t.autofit = False
    for c, text in zip(t.rows[0].cells, headers):
        c.text = text
    for items in rows:
        for c, text in zip(t.add_row().cells, items):
            c.text = text
    for index, row in enumerate(t.rows):
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement('w:cantSplit'))
        for col, c in enumerate(row.cells):
            if widths:
                c.width = Cm(widths[col])
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(1)
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.bold = index == 0
    return t


paragraph(f'DAY {a("Course day", whole)} · DEVPULSE — MONETIZATION ONE-PAGER', True, 13)
paragraph(f'Phạm Thanh Đạt · {a("Student identifier")} · Track {a("Course track",whole)} · {a("Price/source check date")}', size=8)
paragraph(f'Cho CTO nhóm {a("ICP minimum developers", whole)}–{a("ICP maximum developers", whole)} dev: chuẩn bị báo cáo Sprint/tháng có nguồn để giải trình delivery và chi phí AI. Các số vận hành là giả định; chưa có eval/pilot thực đo.', size=8.5)

heading('PRICING')
paragraph('Ngân sách: Engineering Operations/Reporting; CTO đề xuất, CEO/CFO ký theo hạn mức. Job: báo cáo AI-ready có nguồn Jira/GitHub/AI ledger, coverage ≥' + a('Data coverage gate', pct) + ', không lỗi nghiêm trọng; CTO vẫn review để Publish. Khóa workspace/kỳ/loại; không charge retry, sửa lỗi hoặc báo cáo fail.')
paragraph('Hybrid: ' + a('Workspace base fee', usd) + '/workspace/tháng cho kết nối, history, support + Usage/báo cáo từ báo cáo đầu; không có report miễn phí trong gói. Attribution ' + value('3_Value_Metric','B10',whole) + '/' + a('Scorecard maximum',whole) + '; Autonomy ' + value('3_Value_Metric','B18',whole) + '/' + a('Scorecard maximum',whole) + ': chỉ có rule thiết kế, chưa có log/eval; không bán Outcome. HITL A: khách final review/rework; vendor QA vào COGS.')
table(['Chỉ số', 'Giá trị', 'Excel'], [
    ['Cost/Job COGS', value('2_Pricing','B5',usd), 'P!B5'],
    ['Giá sàn', value('2_Pricing','B6',whole) + '×: ' + value('2_Pricing','B7',usd), 'P!B6:B7'],
    ['Giá Usage đề xuất', value('2_Pricing','B19',usd) + '/report', 'P!B19'],
    ['Report Gross Margin', value('2_Pricing','B21',pct), 'P!B21'],
    ['GM mục tiêu / containment tối thiểu', value('2_Pricing','B32',pct) + ' / ' + value('2_Pricing','B33',pct), 'P!B32:B33'],
    ['Containment planning (chưa eval)', value('1_Cost_Job','B10',pct), 'C!B10'],
], [8,5,5.9])
paragraph('Baseline ' + a('Planning workspaces',whole) + ' paying workspace: ' + value('1_Cost_Job','B9',whole) + ' thử / ' + value('1_Cost_Job','B11',whole) + ' AI-ready. Chi phí/tháng: API ' + value('1_Cost_Job','B72',usd) + ', Infra ' + value('1_Cost_Job','B74',usd) + ', HITL ' + value('1_Cost_Job','B76',usd) + ', Retry ' + value('1_Cost_Job','B75',usd) + '; Overhead ' + value('1_Cost_Job','B77',usd) + ' tính riêng. (C!B72:B77)', size=8)
paragraph('Fully loaded Cost/Job ' + a('Fully loaded Cost/Job',usd) + '; khác COGS. All-in ' + a('Reports/customer/month',whole) + ' report = ' + a('ARPU',usd) + '/tháng. Neo ' + a('Monthly value anchor',usd) + ' = ' + a('Reports/customer/month',whole) + ' report × ' + a('Net preparation hours saved/report',whole) + 'h × ' + a('Customer preparation labor',usd) + '/h; trần ' + a('Total value-based monthly ceiling',usd) + ' gồm phí nền. GM <' + a('GM risk threshold',pct) + ' nếu containment <' + a('GM50 containment threshold',pct) + '.', size=8)
paragraph('Benchmark cùng nhóm nhu cầu: Swarmia Standard — Seat, ' + value('3_Value_Metric','C26') + '; LinearB Essentials — Hybrid, ' + value('3_Value_Metric','C27') + '. Nguồn: swarmia.com/pricing/; linearb.io/pricing (V!A26:D27).', size=7.5)

heading('GO-TO-MARKET')
paragraph('Một kênh: PLG qua GitHub App, chưa có listing/partner cam kết. ARPU ' + value('4_Channel_Fit','B5',usd) + '; CAC budget ' + value('4_Channel_Fit','B9',usd) + ' (' + value('4_Channel_Fit','B8',whole) + ' tháng, GM bảo thủ). Rep CAC ' + value('4_Channel_Fit','B22',usd) + ', vượt ' + value('4_Channel_Fit','B23',lambda x:f'{x:.2f}×') + '; ' + value('4_Channel_Fit','B16',lambda x:f'{x:.3f}') + ' deal/AE/ngày không là nút thắt. (G!B5:B23)')
paragraph('PLG forecast: ' + a('PLG acquisition spend',usd) + ' / ' + a('GitHub App installs',whole) + ' installs / ' + a('PLG activated customers',whole) + ' activated / ' + a('PLG paid customers',whole) + ' paid; CAC ' + a('PLG CAC',usd) + '. Có founder time và trial reserve; dừng chi nếu vượt ' + a('PLG operating CAC stop threshold',usd) + '. (A!B28:B35,B79)')
paragraph('Pain Moment: ' + value('5_90Day_Plan','B5') + ' CTO đóng Sprint, ghép delivery/PR/cost ở GitHub/Jira. Nhúng: read-only GitHub App + Jira webhook + AI-ledger import; Slack report card. Chưa triển khai. (M!B5:B8)')
table(['Giai đoạn', 'Kế hoạch; owner: ' + a('Owner')], [
    [value('5_90Day_Plan','B11'), a('Pilot orgs',whole) + ' pilot; xác minh budget, permission, ledger; first report ≤' + a('Trial/activation window',whole) + ' ngày.'],
    [value('5_90Day_Plan','C11'), a('GitHub App installs',whole) + ' installs / ' + a('PLG activated customers',whole) + ' activated / ' + a('90-day paid target',whole) + ' paid; giữ acquisition cap, đo giờ công và retention.'],
    [value('5_90Day_Plan','D11'), 'Chỉ mở ngách tiếp khi đạt floor và GM tại volume thật, eval lower bound và retention; không hứa số khách mới.'],
], [4,14.9])
paragraph('Rủi ro pilot ' + a('Pilot orgs',whole) + ' khách: Cost/Job ' + a('Pilot Cost/Job',usd) + ', floor ' + a('Pilot 3x COGS floor',usd) + ' > price. ≥' + a('Minimum paid orgs at floor3x',whole) + ' serviced paid chỉ là minimum ở planning; vẫn cần eval gate. Không dùng economics của paid cohort cho unpaid chạy liên tục. (A!B70:B77)', size=8)

heading('EVIDENCE PACK')
table(['Tài sản', 'Đã có / còn thiếu', 'Owner · deadline'], [
    ['Eval Results', 'Chưa có kết quả; protocol ' + a('Eval cases',whole) + ' case, tách safety pass và containment.', a('Owner') + ' · ' + a('Eval due')],
    ['Procurement Q&A', 'Bản nháp đã có: hallucination, training, exit/export. Chưa có control proof.', a('Owner') + ' · ' + a('Procurement control verification due')],
    ['Pilot Report', 'Protocol ' + a('Pilot orgs',whole) + ' org / ' + a('Pilot duration',whole) + ' tuần; chưa chạy, chưa có consent.', a('Owner') + ' · ' + a('Pilot report due')],
], [3.4,10.5,5])
paragraph('Reader test: chưa có người lạ đọc, chưa ghi Pass hoặc số câu hỏi. AI review chỉ là kiểm tra nội bộ. Workbook và tài liệu chi tiết có nguồn/giả định; kế hoạch pilot không là kết quả pilot.', size=7.5)
paragraph('Excel: C=1_Cost_Job; P=2_Pricing; V=3_Value_Metric; G=4_Channel_Fit; M=5_90Day_Plan; A=7_Assumptions. Mọi số có mapping trong data/one-pager-trace.json. Giá API và FX đã kiểm tra theo ngày trên workbook.', size=7)

doc.core_properties.title = 'Day22 DevPulse Monetization One-Pager'
doc.core_properties.author = 'Phạm Thanh Đạt'
doc.core_properties.subject = 'Hybrid pricing, PLG and honest evidence gaps'
doc.save(OUT)
(ROOT / 'data/one-pager-trace.json').write_text(json.dumps(trace, ensure_ascii=False, indent=2) + '\n')
print(f'Built DOCX with {len(trace)} trace entries.')
