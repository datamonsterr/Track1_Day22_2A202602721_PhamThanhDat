"""Independently validate financial arithmetic and committed lab artifacts."""
import math


def number(value, name, positive=False):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"{name} must be a finite number")
    if value < 0 or (positive and value == 0):
        raise ValueError(f"{name} must be {'positive' if positive else 'nonnegative'}")
    return value


def calculate_model(cost, price, target_margin=.6):
    required = ("B9 B10 B15 B16 B17 B18 B19 B20 B21 B22 B30 B34 B35 B36 B37 "
                "B41 B42 B46 B50 B51 B52 B53 B59 B68").split()
    for key in required:
        number(cost.get(key), key)
    if cost.get("B6") not in ("A", "B"):
        raise ValueError("B6 must be A or B")
    for key in ("B10", "B46", "B51"):
        if cost[key] > 1:
            raise ValueError(f"{key} must be a fraction in [0, 1]")
    if cost["B30"] not in (0, 1):
        raise ValueError("B30 must be 0 or 1")
    for key in ("B9", "B19"):
        if cost[key] <= 0 or not float(cost[key]).is_integer():
            raise ValueError(f"{key} must be a positive integer")
    if cost["B10"] <= 0:
        raise ValueError("completed jobs must be positive")
    number(price, "price", positive=True)
    number(target_margin, "target_margin")
    if target_margin >= 1:
        raise ValueError("target_margin must be below 1")
    n, rate = cost["B9"], cost["B10"]
    turns = cost["B19"]
    cached_llm = (cost["B20"] * cost["B17"] + (turns - 1) * cost["B20"] * cost["B18"]
        + turns * cost["B21"] * cost["B15"] + turns * cost["B22"] * cost["B16"]) / 1_000_000
    llm = cached_llm * (1 - .5 * cost["B30"])
    speech = cost["B34"] * cost["B35"] + cost["B36"] * cost["B37"] / 1_000_000
    infra = cost["B41"] + cost["B42"] * cost["B35"]
    retry = (llm + speech) * cost["B46"]
    qa = cost["B51"] * cost["B52"] / 60 * cost["B50"]
    escalation = cost["B53"] / 60 * cost["B50"] if cost["B6"] == "B" else 0
    base = llm + speech + infra + retry + qa
    completed = n * rate
    cogs = n * (base + (1 - rate) * escalation)
    unit = cogs / completed
    return {"completed": completed, "cogs": cogs, "cost_per_job": unit,
        "price_floor": 3 * unit, "gross_margin": 1 - unit / price,
        "containment_gm60": (base + escalation) / (price * (1 - target_margin) + escalation),
        "llm": llm, "retry": retry, "qa": n * qa, "escalation": n * (1-rate) * escalation,
        "cost_with_overhead": (cogs + cost["B59"]) / completed}


import ast
import hashlib
import json
import operator
import re
import sys
import zipfile
from datetime import date
from pathlib import Path
from xml.etree import ElementTree

import openpyxl
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import range_boundaries

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = 'deliverables/Day22-DevPulse-Monetization-Model.xlsx'


class FormulaEvaluator:
    """Evaluate the template's small formula language without trusting Excel caches."""
    def __init__(self, book):
        self.book = book
        self.cache = {}
        self.active = set()

    def value(self, sheet, address):
        key = (sheet, address.replace('$', ''))
        if key in self.cache:
            return self.cache[key]
        if key in self.active:
            raise ValueError(f'Cyclic formula {key}')
        self.active.add(key)
        value = self.book[key[0]][key[1]].value
        try:
            if isinstance(value, str) and value.startswith('='):
                expression = []
                for token in Tokenizer(value).items:
                    if token.type == 'OPERAND' and token.subtype == 'RANGE':
                        expression.append(f'REF({token.value!r})')
                    elif token.type == 'OPERAND' and token.subtype == 'TEXT':
                        expression.append(repr(token.value[1:-1].replace('""', '"')))
                    elif token.type == 'OPERAND' and token.subtype == 'LOGICAL':
                        expression.append('True' if token.value == 'TRUE' else 'False')
                    elif token.value == '=':
                        expression.append('==')
                    elif token.value == '<>':
                        expression.append('!=')
                    elif token.value == '^':
                        expression.append('**')
                    else:
                        expression.append(token.value)
                tree = ast.parse(''.join(expression), mode='eval')
                value = self.evaluate(tree.body, sheet)
            self.cache[key] = value
            return value
        finally:
            self.active.remove(key)

    def evaluate(self, node, sheet):
        if isinstance(node, ast.Constant):
            return node.value
        if isinstance(node, ast.BinOp):
            operations = {ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul,
                          ast.Div: operator.truediv, ast.Pow: operator.pow}
            if type(node.op) not in operations:
                raise ValueError('Unsupported arithmetic')
            return operations[type(node.op)](self.evaluate(node.left, sheet), self.evaluate(node.right, sheet))
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.USub, ast.UAdd)):
            value = self.evaluate(node.operand, sheet)
            return -value if isinstance(node.op, ast.USub) else value
        if isinstance(node, ast.Compare) and len(node.ops) == 1:
            operations = {ast.Eq: operator.eq, ast.NotEq: operator.ne, ast.Lt: operator.lt,
                          ast.LtE: operator.le, ast.Gt: operator.gt, ast.GtE: operator.ge}
            return operations[type(node.ops[0])](self.evaluate(node.left, sheet), self.evaluate(node.comparators[0], sheet))
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name):
            function = node.func.id
            if function == 'IF':
                return self.evaluate(node.args[1] if self.evaluate(node.args[0], sheet) else node.args[2], sheet)
            if function == 'REF':
                reference = self.evaluate(node.args[0], sheet).replace('$', '')
                if '!' in reference:
                    target, reference = reference.rsplit('!', 1)
                    sheet = target.strip("'").replace("''", "'")
                if ':' in reference:
                    x1, y1, x2, y2 = range_boundaries(reference)
                    return [self.value(sheet, self.book[sheet].cell(y,x).coordinate)
                            for y in range(y1, y2+1) for x in range(x1,x2+1)]
                return self.value(sheet, reference)
            values = [self.evaluate(arg, sheet) for arg in node.args]
            flat = [item for value in values for item in (value if isinstance(value,list) else [value])]
            if function == 'SUM':
                return sum(item for item in flat if isinstance(item,(int,float)))
            if function == 'MAX':
                return max(flat)
            if function == 'AND':
                return all(flat)
        raise ValueError(f'Unsupported formula expression: {ast.dump(node)}')


def require(condition, message):
    if not condition:
        raise ValueError(message)


def equal(actual, expected, label):
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        require(isinstance(actual, (int,float)) and not isinstance(actual,bool)
                and math.isfinite(actual) and math.isclose(actual, expected, rel_tol=1e-9, abs_tol=1e-9),
                f'{label}: expected {expected!r}, found {actual!r}')
    else:
        require(actual == expected, f'{label}: expected {expected!r}, found {actual!r}')


def read_inputs(root):
    return json.loads((root / 'data/workbook-inputs.json').read_text())


def validate_templates(root):
    manifest = json.loads((root / 'data/template-manifest.json').read_text())
    for entry in manifest:
        path = root / entry['path']
        require(path.is_file(), f'Missing template: {path}')
        require(hashlib.sha256(path.read_bytes()).hexdigest() == entry['sha256'],
                f'Template hash mismatch: {entry["path"]}')
    return len(manifest)


def validate_sources(root):
    inputs = read_inputs(root)
    checked = date.fromisoformat(inputs['checked_on'])
    require(checked <= date.today(), 'Source check date is in the future')
    sources = json.loads((root / 'data/sources.json').read_text())
    by_id = {source['id']:source for source in sources}
    for source in sources:
        require(source['url'].startswith('https://'), f'Missing HTTPS source: {source["id"]}')
        require(date.fromisoformat(source['checked_on']) == checked, f'Source date mismatch: {source["id"]}')
        require(source.get('scope') and source.get('status'), f'Missing source scope/status: {source["id"]}')
    cost = inputs['sheets']['1_Cost_Job']
    prices = by_id['anthropic-pricing']['facts']
    for address, fact in {'B15':'input_per_mtok','B16':'output_per_mtok',
                          'B17':'cache_write_5m_per_mtok','B18':'cache_read_per_mtok'}.items():
        equal(cost[address], prices[fact], f'Haiku price {address}')
    minimum = by_id['anthropic-cache']['facts']['minimum_cache_tokens_haiku45']
    require(cost['B20'] >= minimum, f'Haiku cache prefix below minimum {minimum}')
    require('5' in inputs['notes']['1_Cost_Job!B20'] or 'five' in inputs['notes']['1_Cost_Job!B20'],
            'Missing cache TTL assumption')
    return len(sources)


def validate_workbook(root):
    inputs = read_inputs(root)
    path = root / WORKBOOK
    require(path.is_file(), f'Missing workbook: {path}')
    original = openpyxl.load_workbook(root / 'templates/Day22-AI-Product-GTM-Monetization-Model.xlsx')
    book = openpyxl.load_workbook(path)
    cached = openpyxl.load_workbook(path, data_only=True)
    evaluator = FormulaEvaluator(book)
    count = 0
    for sheet in original:
        require(sheet.title in book.sheetnames, f'Missing original sheet {sheet.title}')
        submitted = inputs['sheets'].get(sheet.title,{})
        for row in sheet:
            for cell in row:
                target = book[sheet.title][cell.coordinate]
                yellow = cell.fill.fgColor.type == 'rgb' and cell.fill.fgColor.rgb == 'FFFFF2CC'
                if cell.data_type == 'f' or not yellow:
                    equal(target.value, cell.value, f'Original cell {sheet.title}!{cell.coordinate}')
                elif sheet.title != '0_README':
                    if sheet.title == '6_Benchmarks':
                        expected = inputs['checked_on']
                    else:
                        # Earlier checkpoints may deliberately lack later station inputs.
                        required_stage = {'1_Cost_Job':1,'2_Pricing':3,'3_Value_Metric':2,
                                          '4_Channel_Fit':4,'5_90Day_Plan':5}[sheet.title]
                        if inputs['stage'] < required_stage:
                            continue
                        require(cell.coordinate in submitted, f'Missing yellow input {sheet.title}!{cell.coordinate}')
                        expected = submitted[cell.coordinate]
                    require(expected is not None and expected != '', f'Empty yellow input {sheet.title}!{cell.coordinate}')
                    equal(target.value, expected, f'Yellow input {sheet.title}!{cell.coordinate}')
    for sheet in book:
        for row in sheet:
            for cell in row:
                if cell.data_type == 'f':
                    expected = evaluator.value(sheet.title, cell.coordinate)
                    equal(cached[sheet.title][cell.coordinate].value, expected,
                          f'Formula cache {sheet.title}!{cell.coordinate}')
                    count += 1
    model = calculate_model(inputs['sheets']['1_Cost_Job'], inputs['sheets']['2_Pricing']['B19'])
    for sheet, address, key in [('1_Cost_Job','B11','completed'),('1_Cost_Job','B66','cost_per_job'),
                              ('1_Cost_Job','B67','cost_with_overhead'),('2_Pricing','B7','price_floor'),
                              ('2_Pricing','B21','gross_margin'),('2_Pricing','B33','containment_gm60')]:
        equal(cached[sheet][address].value, model[key], f'Independent math {sheet}!{address}')
    require(model['gross_margin'] >= .6, 'Gross margin is below 60%')
    require(model['gross_margin'] <= .85, 'Gross margin exceeds 85%; cost audit required')
    require(inputs['sheets']['2_Pricing']['B19'] >= model['price_floor'], 'Price below 3x Cost/Job floor')
    return count


def validate_formatted_number(value, formatted, label):
    """Require a displayed numerical claim to round the traced workbook value."""
    require(not isinstance(value, bool) and isinstance(value, (int, float)) and math.isfinite(value),
            f'Formatted value source must be finite: {label}')
    tokens = re.findall(r"-?\d+(?:[.,]\d+)*", str(formatted))
    require(bool(tokens), f'Formatted value has no number: {label}')
    token = tokens[0]
    target = value * 100 if '%' in str(formatted) else value
    candidates = []
    if ',' in token and '.' in token:
        decimal = ',' if token.rfind(',') > token.rfind('.') else '.'
        grouping = '.' if decimal == ',' else ','
        normalized = token.replace(grouping, '').replace(decimal, '.')
        candidates.append((float(normalized), len(token.rsplit(decimal, 1)[1])))
    else:
        separator = ',' if ',' in token else '.' if '.' in token else None
        if separator:
            digits = len(token.rsplit(separator, 1)[1])
            candidates.append((float(token.replace(separator, '.')), digits))
            if digits == 3:
                candidates.append((float(token.replace(separator, '')), 0))
        else:
            candidates.append((float(token), 0))
    require(any(abs(shown-target) <= .5*10**(-precision)+1e-10 for shown,precision in candidates),
            f'Formatted value disagrees with workbook: {label}: {formatted!r} vs {value!r}')


def numeric_claims(text, sheet_names):
    """Exclude URLs and precise workbook identifiers, retaining all numeric claims."""
    claims = re.sub(r"https?://\S+", "", text)
    cell = r"\$?[A-Z]+\$?\d+"
    references = rf"(?:'[^']+'|[A-Za-z0-9_]+)!{cell}(?::{cell})?(?:,{cell}(?::{cell})?)*"
    claims = re.sub(references, "", claims)
    for sheet_name in sheet_names:
        claims = claims.replace(sheet_name, '')
    return set(re.findall(r"\d+(?:[.,]\d+)*", claims))


def validate_final_artifacts(root):
    if read_inputs(root)['stage'] < 6:
        return {'status':'not required before checkpoint 6'}
    documents = list((root / 'deliverables').glob('*One-Pager*.docx'))
    pdfs = list((root / 'deliverables').glob('*One-Pager*.pdf'))
    trace_path = root / 'data/one-pager-trace.json'
    require(len(documents) == 1 and len(pdfs) == 1 and trace_path.is_file(), 'Missing final DOCX/PDF/source-map artifacts')
    import fitz
    from docx import Document
    document = Document(documents[0])
    text = '\n'.join([p.text for p in document.paragraphs] +
                     [c.text for table in document.tables for row in table.rows for c in row.cells])
    with zipfile.ZipFile(documents[0]) as archive:
        app = ElementTree.fromstring(archive.read('docProps/app.xml'))
        pages = app.find('{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}Pages')
        require(pages is not None and pages.text == '1', 'DOCX must record exactly one rendered page')
        xml = ElementTree.fromstring(archive.read('word/document.xml'))
        namespace = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        require(not any(el.get(namespace+'type') == 'page' for el in xml.iter(namespace+'br')),
                'DOCX contains explicit page break')
    with fitz.open(pdfs[0]) as pdf:
        require(len(pdf) == 1, f'PDF must be exactly one page, found {len(pdf)}')
        pdf_text = pdf[0].get_text()
    trace = json.loads(trace_path.read_text())
    require(isinstance(trace,list) and bool(trace), 'Source map must be a nonempty list')
    cached = openpyxl.load_workbook(root / WORKBOOK, data_only=True)
    for entry in trace:
        require(set(('label','sheet','cell','value','formatted')) <= entry.keys(), 'Incomplete source-map entry')
        require(entry['sheet'] in cached.sheetnames, f'Unknown source-map sheet {entry["sheet"]}')
        equal(entry['value'], cached[entry['sheet']][entry['cell']].value, f'Source map {entry["label"]}')
        if isinstance(entry['value'], (int, float)) and not isinstance(entry['value'], bool):
            validate_formatted_number(entry['value'], entry['formatted'], entry['label'])
        require(str(entry['formatted']) in text, f'Source-map formatted value absent in DOCX: {entry["label"]}')
    # Every numeric token visible in the DOCX must come from a traced formatted value.
    numeric = re.compile(r'\d+(?:[.,]\d+)*')
    claims = numeric_claims(text, cached.sheetnames)
    allowed = {token for entry in trace for token in numeric.findall(str(entry['formatted']))}
    require(claims <= allowed, f'Untraced DOCX numbers: {sorted(claims - allowed)}')
    # Text extracted from PDF must retain every substantive document word.
    words = lambda value: set(re.findall(r'[^\W\d_]+', value.lower(), re.UNICODE))
    require(words(text) <= words(pdf_text), f'DOCX content missing from PDF: {sorted(words(text)-words(pdf_text))[:8]}')
    return {'pages':1,'trace_entries':len(trace)}


def main():
    try:
        result = {'template_hashes':validate_templates(ROOT), 'sources':validate_sources(ROOT),
                  'formula_caches':validate_workbook(ROOT), 'final':validate_final_artifacts(ROOT)}
        print(json.dumps(result, ensure_ascii=False))
    except (ValueError, KeyError, TypeError, OSError) as error:
        print(f'Validation failed: {error}', file=sys.stderr)
        return 1
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
