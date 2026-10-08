"""Integration and contract tests for the committed submission, never rebuilt in CI."""
import json
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts.validate_model import FormulaEvaluator, validate_workbook, validate_sources, validate_final_artifacts, validate_templates


def test_excel_evaluator_handles_cross_sheet_ranges_and_lazy_guards():
    book = openpyxl.Workbook()
    first = book.active
    first.title = "one"
    second = book.create_sheet("two")
    first["A1"] = 2
    first["A2"] = 3
    first["B1"] = '=SUM(A1:A2)*2'
    second["A1"] = "='one'!B1/2"
    second["A2"] = '=IF(AND(A1=5,A1>=3),"OK",1/0)'
    evaluator = FormulaEvaluator(book)
    assert evaluator.value("two", "A1") == 5
    assert evaluator.value("two", "A2") == "OK"


def test_excel_evaluator_rejects_unknown_function():
    book = openpyxl.Workbook()
    book.active["A1"] = '=UNSUPPORTED(1)'
    with pytest.raises(ValueError, match="Unsupported"):
        FormulaEvaluator(book).value(book.active.title,"A1")


def test_original_templates_match_manifest():
    validate_templates(ROOT)


def test_sources_and_haiku_cache_constraints():
    validate_sources(ROOT)


def test_workbook_integrity_yellow_inputs_and_formula_caches():
    validate_workbook(ROOT)


def test_final_onepage_and_trace_contract():
    stage = json.loads((ROOT / "data/workbook-inputs.json").read_text())["stage"]
    if stage < 6:
        pytest.skip("Final DOCX/PDF/trace are required at checkpoint 6 only")
    validate_final_artifacts(ROOT)


def test_stage6_does_not_silently_skip_missing_artifacts(tmp_path):
    (tmp_path / "data").mkdir()
    (tmp_path / "data/workbook-inputs.json").write_text('{"stage": 6}')
    with pytest.raises(ValueError, match="Missing"):
        validate_final_artifacts(tmp_path)



@pytest.fixture
def copied_inputs(tmp_path):
    import shutil
    shutil.copytree(ROOT / "data", tmp_path / "data")
    return tmp_path


def test_price_provenance_mismatch_is_rejected(copied_inputs):
    path = copied_inputs / "data/workbook-inputs.json"
    inputs = json.loads(path.read_text())
    inputs["sheets"]["1_Cost_Job"]["B15"] = 2
    path.write_text(json.dumps(inputs))
    with pytest.raises(ValueError, match="Haiku price"):
        validate_sources(copied_inputs)


def test_ineligible_haiku_cache_is_rejected(copied_inputs):
    path = copied_inputs / "data/workbook-inputs.json"
    inputs = json.loads(path.read_text())
    inputs["sheets"]["1_Cost_Job"]["B20"] = 3000
    path.write_text(json.dumps(inputs))
    with pytest.raises(ValueError, match="below minimum"):
        validate_sources(copied_inputs)


def test_mismatched_source_dates_are_rejected(copied_inputs):
    path = copied_inputs / "data/sources.json"
    sources = json.loads(path.read_text())
    sources[0]["checked_on"] = "2026-10-07"
    path.write_text(json.dumps(sources))
    with pytest.raises(ValueError, match="Source date mismatch"):
        validate_sources(copied_inputs)


def test_modified_original_template_is_rejected(tmp_path):
    import shutil
    shutil.copytree(ROOT / "templates", tmp_path / "templates")
    (tmp_path / "data").mkdir()
    shutil.copy(ROOT / "data/template-manifest.json", tmp_path / "data")
    template = tmp_path / "templates/Day22-AI-Product-GTM-Monetization-Model.xlsx"
    template.write_bytes(template.read_bytes() + b"corruption")
    with pytest.raises(ValueError, match="Template hash mismatch"):
        validate_templates(tmp_path)


@pytest.fixture
def copied_workbook(copied_inputs):
    import shutil
    shutil.copytree(ROOT / 'templates', copied_inputs / 'templates')
    (copied_inputs / 'deliverables').mkdir()
    shutil.copy(ROOT / 'deliverables/Day22-DevPulse-Monetization-Model.xlsx', copied_inputs / 'deliverables')
    return copied_inputs


@pytest.mark.parametrize('sheet,address,value', [
    ('1_Cost_Job','B11','=B9'),
    ('2_Pricing','A39',.1),
])
def test_non_yellow_values_and_formulas_are_immutable(copied_workbook, sheet, address, value):
    path = copied_workbook / 'deliverables/Day22-DevPulse-Monetization-Model.xlsx'
    book = openpyxl.load_workbook(path)
    book[sheet][address] = value
    book.save(path)
    with pytest.raises(ValueError, match='Original cell'):
        validate_workbook(copied_workbook)


def test_missing_yellow_working_input_is_rejected(copied_workbook):
    path = copied_workbook / 'data/workbook-inputs.json'
    data = json.loads(path.read_text())
    del data['sheets']['1_Cost_Job']['B41']
    path.write_text(json.dumps(data))
    with pytest.raises(ValueError, match='Missing yellow input'):
        validate_workbook(copied_workbook)


def test_excel_evaluator_detects_cyclic_formulas():
    book = openpyxl.Workbook()
    book.active['A1'] = '=B1'
    book.active['B1'] = '=A1'
    with pytest.raises(ValueError, match='Cyclic formula'):
        FormulaEvaluator(book).value(book.active.title, 'A1')


def test_missing_formula_cache_is_rejected(copied_workbook):
    import zipfile
    from xml.etree import ElementTree
    path = copied_workbook / 'deliverables/Day22-DevPulse-Monetization-Model.xlsx'
    with zipfile.ZipFile(path) as archive:
        files = {name:archive.read(name) for name in archive.namelist()}
    xml = ElementTree.fromstring(files['xl/worksheets/sheet2.xml'])
    namespace = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    target = next(cell for cell in xml.iter(namespace+'c') if cell.get('r') == 'B11')
    value = target.find(namespace+'v')
    if value is not None:
        target.remove(value)
    files['xl/worksheets/sheet2.xml'] = ElementTree.tostring(xml)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    with pytest.raises(ValueError, match='Formula cache 1_Cost_Job!B11'):
        validate_workbook(copied_workbook)
